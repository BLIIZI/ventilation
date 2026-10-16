from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

F = 'Arial'
BLUE = Font(name=F, color='0000FF'); BLK = Font(name=F); B = Font(name=F, bold=True)
H = Font(name=F, bold=True, color='FFFFFF'); HFILL = PatternFill('solid', fgColor='305496')
YEL = PatternFill('solid', fgColor='FFF2CC'); GRN = PatternFill('solid', fgColor='E2EFDA'); RED = PatternFill('solid', fgColor='FCE4D6')
thin = Side(style='thin', color='999999'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

def hdr(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = HFILL; c.alignment = WR; c.border = BOX

def setw(ws, widths):
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w

# ---------------- Легенда
ws = wb.active; ws.title = 'Легенда'
rows = [
    ('Блок модели «вещества → потребный воздухообмен». Версия 1 от 09.10.2026', B),
    ('', BLK),
    ('Как читать файл', B),
    ('Синий шрифт на жёлтом фоне — исходные данные, их можно менять.', BLK),
    ('Чёрный шрифт — формулы (видны в строке формул), пересчитываются автоматически.', BLK),
    ('Зелёный фон — итоговый результат.', BLK),
    ('Красный фон — расхождение с дипломом или значение, которое нужно проверить.', BLK),
    ('', BLK),
    ('Формулы модели', B),
    ('(1) G_i = g_i · 1000 · B · n — выделение вещества i, мг/ч (g_i в г/кг, B — расход сварочного материала на пост, кг/ч, n — число постов).', BLK),
    ('     Источник: Методика расчёта выделений при сварочных работах на основе удельных показателей (прил. 2 к приказу Госкомэкологии РФ № 158 от 14.04.1997); ГОСТ Р 56164-2014.', BLK),
    ('(2) G_пом,i = (1 − η) · G_i — часть, не уловленная местным отсосом, мг/ч. η — эффективность захвата: ДОПУЩЕНИЕ (в дипломе 0,75).', BLK),
    ('(3) L_i = G_пом,i / (ПДК_i − q_пр,i) — общеобменный воздухообмен по веществу i, м³/ч. Формула расчёта воздухообмена по массе вредных веществ СП 60.13330 (как в дипломе, при L_wz = 0 и m_po = G_пом).', BLK),
    ('     q_пр — концентрация в приточном воздухе. Для наружного воздуха — фон, сверху ограничен ПДК атм. воздуха (как в дипломе). Ограничение: q_пр ≤ 0,3·ПДК р.з. (СП 60.13330.2020, п. 5.11).', BLK),
    ('(4) Несколько веществ: расход по веществу, требующему наибольшего воздухообмена; для веществ однонаправленного действия расходы суммируются (СП 60.13330.2020, п. 7.4.3).', BLK),
    ('     В таблице это реализовано так: вещества одной «группы» суммируются, итог = максимум по группам.', BLK),
    ('(5) Местный отсос: L_мо ≥ 1500 м³/ч на стационарный пост со сварочным столом, скорость у точки сварки ≥ 0,2 м/с (ГОСТ 12.3.003-86, п. 2.11.1).', BLK),
    ('', BLK),
    ('Что нужно проверить по первоисточникам', B),
    ('— ПДК фтористого водорода (0,1) и ПДК атм. воздуха — взяты из диплома, сверить с СанПиН 1.2.3685-21.', BLK),
    ('— Удельные выделения флюса ФЦ-12 — взяты из диплома (табл. 1.6), сверить с Методикой 1997 г. Значения для Св-08Г2С подтверждены вторым источником.', BLK),
    ('— ПДК марганца зависит от его доли в аэрозоле (до 20 % — 0,6/0,2; 20–30 % — 0,3/0,1). У ФЦ-12 доля Mn = 33 % — категорию нужно уточнить по СанПиН.', BLK),
    ('— Однонаправленность действия оксида железа и HF (принята в дипломе) требует обоснования.', BLK),
]
for i, (t, f) in enumerate(rows, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = f
ws.column_dimensions['A'].width = 160

# ---------------- Справочник
sp = wb.create_sheet('Справочник')
hdr(sp, 1, ['Вещество', 'ПДК р.з. ср.см, мг/м³', 'ПДК атм. ср.сут, мг/м³ (фон q_пр)', 'Статус', 'Источник'])
subs = [
    ('Железа оксид', 6, 0.04, 'Подтверждено вторым источником', 'Диплом, табл. 2.2; Быстряков, Невар // Веснік ВДУ, 2020, № 1(106) — 6,0 мг/м³'),
    ('Марганец в сварочных аэрозолях (Mn ≤ 20 %)', 0.2, 0.001, 'Подтверждено', 'Походня И.К. и др. // Автоматическая сварка, 2011, № 9 — 0,6/0,2 мг/м³; ПДК атм. — диплом'),
    ('Фтористый водород', 0.1, 0.01, 'Проверить', 'Диплом, табл. 2.2; сверить с СанПиН 1.2.3685-21'),
    ('Углерода оксид', 20, None, 'Подтверждено вторым источником', 'Быстряков, Невар, 2020 — 20 мг/м³'),
    ('Азота диоксид', 2, None, 'Подтверждено вторым источником (макс.)', 'Быстряков, Невар, 2020 — 2,0 мг/м³ (максимально разовая)'),
]
for r, row in enumerate(subs, 2):
    for c, v in enumerate(row, 1):
        cell = sp.cell(row=r, column=c, value=v); cell.border = BOX; cell.alignment = WR
        cell.font = BLUE if c in (2, 3) else BLK
        if c in (2, 3) and v is not None: cell.fill = YEL
    if row[3] == 'Проверить':
        sp.cell(row=r, column=4).fill = RED
r0 = len(subs) + 4
sp.cell(row=r0, column=1, value='Удельные выделения, г/кг расходуемого сварочного материала').font = B
hdr(sp, r0 + 1, ['Материал / процесс', 'Сварочный аэрозоль', 'Железа оксид', 'Марганец', 'Пыль SiO₂ 20–70 %', 'Фтористый водород', 'Статус / источник'])
mats = [
    ('Св-08Г2С, сварка в CO₂', 10.00, 7.67, 1.90, 0.43, None, 'Методика 1997 г. (по диплому, табл. 1.6); 7,67 / 1,9 / 0,43 подтверждены методикой РК (ecoportal.kz)'),
    ('ФЦ-12, сварка под флюсом', 0.09, 0.06, 0.03, None, 0.02, 'ПРОВЕРИТЬ: диплом, табл. 1.6; сверить с Методикой 1997 г.'),
]
for i, row in enumerate(mats):
    r = r0 + 2 + i
    for c, v in enumerate(row, 1):
        cell = sp.cell(row=r, column=c, value=v); cell.border = BOX; cell.alignment = WR
        cell.font = BLUE if 2 <= c <= 6 else BLK
        if 2 <= c <= 6 and v is not None: cell.fill = YEL
    if 'ПРОВЕРИТЬ' in row[6]: sp.cell(row=r, column=7).fill = RED
SP_SW = r0 + 2; SP_FC = r0 + 3  # строки материалов
for col, w in zip('ABCDEFG', [42, 16, 18, 14, 16, 16, 70]): sp.column_dimensions[col].width = w

# ---------------- расчётный лист
def calc_sheet(name, title, B_val, n_val, eta_val, subs_rows, diploma=None, note=''):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = B
    ws['A2'] = note; ws['A2'].font = BLK
    inputs = [('Расход сварочного материала на пост B, кг/ч', B_val),
              ('Число постов n', n_val),
              ('Эффективность захвата местным отсосом η (допущение)', eta_val),
              ('Местный отсос на пост L_мо, м³/ч (ГОСТ 12.3.003-86, п. 2.11.1: ≥ 1500)', 1500)]
    for i, (t, v) in enumerate(inputs, 4):
        ws.cell(row=i, column=1, value=t).font = BLK
        c = ws.cell(row=i, column=2, value=v); c.font = BLUE; c.fill = YEL; c.border = BOX
    ws['B6'].number_format = '0%'
    hdr(ws, 9, ['Вещество', 'g, г/кг', 'G = g·1000·B·n, мг/ч', 'G_пом = (1−η)·G, мг/ч', 'ПДК р.з., мг/м³', 'q_пр, мг/м³',
                'L = G_пом/(ПДК−q_пр), м³/ч', 'Группа однонаправл. действия'] + (['Диплом, м³/ч', 'Расхождение'] if diploma else []))
    first = 10
    for i, (label, gref, pdkref, qref, grp) in enumerate(subs_rows):
        r = first + i
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=f'={gref}')
        ws.cell(row=r, column=3, value=f'=B{r}*1000*$B$4*$B$5')
        ws.cell(row=r, column=4, value=f'=(1-$B$6)*C{r}')
        ws.cell(row=r, column=5, value=f'={pdkref}')
        ws.cell(row=r, column=6, value=f'={qref}')
        ws.cell(row=r, column=7, value=f'=IF(E{r}>F{r},D{r}/(E{r}-F{r}),"ПДК ≤ фона")')
        g = ws.cell(row=r, column=8, value=grp); g.font = BLUE; g.fill = YEL
        for c in range(2, 7): ws.cell(row=r, column=c).font = Font(name=F, color='008000') if c in (2, 5, 6) else BLK
        if diploma:
            d = ws.cell(row=r, column=9, value=diploma[i]); d.font = BLUE
            ws.cell(row=r, column=10, value=f'=IF(ABS(G{r}-I{r})>2,G{r}-I{r},"совпадает")')
        for c in range(1, 11 if diploma else 9): ws.cell(row=r, column=c).border = BOX
        for c in (3, 4, 7): ws.cell(row=r, column=c).number_format = '#,##0.0'
    last = first + len(subs_rows) - 1
    R = last + 2
    ws.cell(row=R, column=1, value='Сумма L по группам').font = B
    for k in range(1, 4):
        ws.cell(row=R + k, column=1, value=f'Группа {k}')
        c = ws.cell(row=R + k, column=7, value=f'=SUMIF($H${first}:$H${last},{k},$G${first}:$G${last})'); c.number_format = '#,##0.0'
    res = R + 5
    rows = [
        ('Общеобменный воздухообмен L_оо = max по группам, м³/ч (СП 60.13330.2020, п. 7.4.3)', f'=MAX(G{R+1}:G{R+3})'),
        ('Для сравнения: если все вещества разного действия (максимум), м³/ч', f'=MAX(G{first}:G{last})'),
        ('Местные отсосы, всего ΣL_мо = L_мо·n, м³/ч', '=B7*B5'),
        ('Итого удаляемый (= приточный) воздух L_мо + L_оо, м³/ч', f'=G{res}+G{res+2}'),
    ]
    for j, (t, f) in enumerate(rows):
        ws.cell(row=res + j, column=1, value=t).font = B if j in (0, 3) else BLK
        c = ws.cell(row=res + j, column=7, value=f); c.number_format = '#,##0.0'; c.border = BOX
        if j in (0, 3): c.fill = GRN; c.font = B
    for col, w in zip('ABCDEFGHIJ', [62, 10, 18, 18, 14, 12, 20, 16, 14, 14]): ws.column_dimensions[col].width = w
    ws.row_dimensions[9].height = 45
    return ws, first, last, res

S = "'Справочник'!"
dip_subs = [('Железа оксид', f'{S}C{SP_FC}', f'{S}B2', f'{S}C2', 1),
            ('Марганец', f'{S}D{SP_FC}', f'{S}B3', f'{S}C3', 2),
            ('Фтористый водород', f'{S}F{SP_FC}', f'{S}B4', f'{S}C4', 1)]
ws1, f1, l1, res1 = calc_sheet('Диплом_под_флюсом', 'Контрольный пример 1: диплом, 4 поста, автоматическая сварка под флюсом ФЦ-12',
                               6.808, 4, 0.75, dip_subs, diploma=[68, 1075, 1514],
                               note='Группы как в дипломе: оксид железа и HF — группа 1 (суммируются), Mn — группа 2. Чтобы считать «каждое отдельно», поставь HF группу 3.')
ws1['L10'] = 'Проверка исходной строки диплома без захвата: 817 / (0,2 − 0,001) ='
ws1['L11'] = '=817/(0.2-0.001)'; ws1['L11'].number_format = '#,##0.0'
ws1['L12'] = 'В дипломе указано 4300 м³/ч — арифметическая ошибка'; ws1['L12'].fill = RED
ws1.column_dimensions['L'].width = 60
ws1.cell(row=f1 + 1, column=10).fill = RED

mag_subs = [('Железа оксид', f'{S}C{SP_SW}', f'{S}B2', f'{S}C2', 1),
            ('Марганец', f'{S}D{SP_SW}', f'{S}B3', f'{S}C3', 2)]
ws2, f2, l2, res2 = calc_sheet('MAG_в_CO2', 'Пример 2: полуавтоматическая сварка в CO₂ проволокой Св-08Г2С, 1 пост',
                               3.0, 1, 0.75, mag_subs,
                               note='B = 3 кг/ч — условное значение для примера (не из источника). Пыль SiO₂ не включена: ПДК не проверена.')
# чувствительность к η
r = res2 + 6
ws2.cell(row=r, column=1, value='Чувствительность к эффективности захвата η').font = B
hdr(ws2, r + 1, ['η', 'L по Mn, м³/ч', 'L по Fe, м³/ч', 'L_оо, м³/ч', 'Итого с местным, м³/ч'])
for i, eta in enumerate([0.6, 0.75, 0.9]):
    rr = r + 2 + i
    c = ws2.cell(row=rr, column=1, value=eta); c.font = BLUE; c.fill = YEL; c.number_format = '0%'
    ws2.cell(row=rr, column=2, value=f'=(1-A{rr})*C{f2+1}/(E{f2+1}-F{f2+1})')
    ws2.cell(row=rr, column=3, value=f'=(1-A{rr})*C{f2}/(E{f2}-F{f2})')
    ws2.cell(row=rr, column=4, value=f'=MAX(B{rr}:C{rr})')
    ws2.cell(row=rr, column=5, value=f'=D{rr}+$B$7*$B$5')
    for c in range(1, 6):
        ws2.cell(row=rr, column=c).border = BOX
        if c > 1: ws2.cell(row=rr, column=c).number_format = '#,##0'

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, color=c.font.color)
wb.save('/home/claude/model/Модель_воздухообмен_v1.xlsx')
print('ok')
